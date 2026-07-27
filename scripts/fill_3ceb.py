"""
fill_3ceb.py

Writes extracted/reviewed Part A data into a 3CEB Excel template.

CELL_MAP below is a PLACEHOLDER. Open your actual blank 3CEB template, find the sheet name
and cell address for each Clause 1-9 field, and update CELL_MAP accordingly before using this
for real. Do not guess these values against a real filing.

Usage:
    python fill_3ceb.py --data reviewed.json --template blank_3ceb.xlsx --out draft_3ceb.xlsx
"""

import argparse
import json

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

FLAG_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

# --- ADJUST THESE to match your actual template's sheet/cell layout ---
SHEET_NAME = "Part A"
CELL_MAP = {
    "assessee_name": "C4",
    "pan": "C5",
    "address": "C6",
    "business_code": "C7",
    "aggregate_international_transactions": "C8",
    "aggregate_sdt": "C9",
}
# ------------------------------------------------------------------------


def fill(data_path, template_path, out_path):
    with open(data_path) as f:
        data = json.load(f)

    wb = load_workbook(template_path)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(
            f"Sheet '{SHEET_NAME}' not found in template. Available sheets: {wb.sheetnames}. "
            f"Update SHEET_NAME in this script."
        )
    ws = wb[SHEET_NAME]

    unfilled = []
    for field, cell_ref in CELL_MAP.items():
        value = data.get(field)
        if value is None:
            unfilled.append(field)
            continue
        ws[cell_ref] = value
        # Flag every auto-filled cell visually so the reviewer knows it wasn't manually entered
        ws[cell_ref].fill = FLAG_FILL

    wb.save(out_path)

    print(f"Saved draft to {out_path}. Auto-filled cells are highlighted for reviewer visibility.")
    if unfilled:
        print(f"NOT filled (missing from extracted data, needs manual entry): {unfilled}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    fill(args.data, args.template, args.out)


if __name__ == "__main__":
    main()
