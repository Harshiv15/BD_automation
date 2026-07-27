# Streamlit Web App for Bulk BD Automation
# To run this locally:
# 1. Install dependencies: pip install streamlit openpyxl google-genai pydantic pdfplumber
# 2. Run command: streamlit run app.py

import streamlit as st
import os
import json
import tempfile
import pdfplumber
from pydantic import BaseModel, Field
from typing import List, Optional

from google import genai
from google.genai import types

try:
    import fill_factsheet
    import safe_insert_rows
except ImportError as e:
    st.error(f"Failed to import required scripts. Error: {e}")
    st.stop()

st.set_page_config(page_title="EY TP BD Automator", layout="wide")

# =====================================================================
# 1. DEFINE THE STRICT JSON SCHEMA (PYDANTIC)
# =====================================================================

class FieldBase(BaseModel):
    value: str | int | float | None = Field(description="The extracted value. Use None if not found.")
    source: str = Field(description="Exact page number and section where this was found.")
    flag: Optional[str] = Field(None, description="Any judgment calls, assumptions, or notes for the human reviewer.")

class AERevenueSplitValue(BaseModel):
    domestic_ae_fy25: float | int | None
    export_ae_fy25: float | int | None
    domestic_third_party_fy25: float | int | None
    export_third_party_fy25: float | int | None

class AERevenueSplit(BaseModel):
    value: AERevenueSplitValue
    source: str
    flag: Optional[str] = None

class ShareholderRow(BaseModel):
    name: str = Field(description="Name of the shareholder/promoter category")
    shares_fy25: int | float | None
    shares_fy24: int | float | None

class Shareholding(BaseModel):
    rows: List[ShareholderRow]
    total_shares_fy25: int | float | None
    total_shares_fy24: int | float | None
    source: str

class RPTItem(BaseModel):
    label: str = Field(description="Exact transaction label from the AR")
    value_fy25: float | int | None = Field(description="Positive for outflow/cost, Negative for inflow/receivable")

class RelatedPartyTransactions(BaseModel):
    items: List[RPTItem]
    source: str

class LitigationItem(BaseModel):
    nature_of_dues: str | None
    amount_demanded_lakhs: float | int | None
    amount_paid_lakhs: float | int | None
    period: str | None
    forum: str | None

class Litigation(BaseModel):
    items: List[LitigationItem]
    source: str

class ExtractedFields(BaseModel):
    hq_india_entity: FieldBase
    hq_group: FieldBase
    company_description: FieldBase
    group_description: FieldBase
    standalone_turnover_fy25_lakhs: FieldBase
    standalone_turnover_fy24_lakhs: FieldBase
    standalone_total_cost_fy25_lakhs: FieldBase
    standalone_total_cost_fy24_lakhs: FieldBase
    consolidated_summary: FieldBase
    statutory_auditors: FieldBase
    ae_revenue_split: AERevenueSplit
    cash_fy25_lakhs: FieldBase
    cash_fy24_lakhs: FieldBase
    ae_trade_receivables_fy25_lakhs: FieldBase
    ae_trade_receivables_fy24_lakhs: FieldBase
    pe_investment: FieldBase
    shareholding: Shareholding
    related_party_transactions_lakhs: RelatedPartyTransactions
    countries_presence: FieldBase
    litigation: Litigation
    website: FieldBase
    linkedin: FieldBase

class CompanyData(BaseModel):
    entity: str = Field(description="Name of the company")
    fields: ExtractedFields


# =====================================================================
# 2. THE DUAL-EXTRACTION ENGINE (TEXT FIRST, VISION BACKUP)
# =====================================================================

def extract_bd_data(pdf_path: str, api_key: str, company_name: str, model_name: str) -> dict:
    client = genai.Client(api_key=api_key)
    
    # Try text extraction first
    extracted_text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_text += text + "\n"
    
    with open("SKILL.md", "r") as f:
        system_prompt = f.read()

    config = types.GenerateContentConfig(
        system_instruction=system_prompt,
        temperature=0.0,
        response_mime_type="application/json",
        response_schema=CompanyData,
    )

    # HEURISTIC FALLBACK: If < 2000 chars, it's a scanned image. Use File API (Vision).
    if len(extracted_text.strip()) < 2000:
        st.info(f"[{company_name}] Detected scanned/image PDF. Falling back to Gemini Vision API (Costs more quota)...")
        uploaded_file = client.files.upload(file=pdf_path)
        user_prompt = f"Please extract the Business Description data for {company_name} from the attached Annual Report PDF."
        
        response = client.models.generate_content(
            model=model_name,
            contents=[uploaded_file, user_prompt],
            config=config
        )
        client.files.delete(name=uploaded_file.name) # cleanup
        
    else:
        st.info(f"[{company_name}] Text successfully extracted ({len(extracted_text)} chars). Using Text API (Saves quota)...")
        user_prompt = f"Please extract the Business Description data for {company_name} from the following Annual Report text:\n\n{extracted_text}"
        
        response = client.models.generate_content(
            model=model_name, 
            contents=user_prompt,
            config=config
        )

    return json.loads(response.text)

# =====================================================================
# 3. STREAMLIT UI 
# =====================================================================

st.title("TP BD Automator")

with st.sidebar:
    st.header("⚙️ Settings")
    api_key = st.text_input("Gemini API Key", type="password")
    model_choice = st.text_input("Model", value="gemini-flash-latest")

col1, col2 = st.columns(2)
with col1:
    excel_template = st.file_uploader("Upload BD Format", type=["xlsx"])
with col2:
    pdf_files = st.file_uploader("Upload ARs", type=["pdf"], accept_multiple_files=True)

if st.button("Start Bulk Extraction", type="primary"):
    if not api_key:
        st.error("Please enter a Gemini API key.")
    elif not excel_template or not pdf_files:
        st.error("Please upload the Excel template and at least one PDF.")
    else:
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = os.path.join(temp_dir, "template.xlsx")
            with open(template_path, "wb") as f:
                f.write(excel_template.getvalue())

            from openpyxl import load_workbook
            
            for i, pdf_file in enumerate(pdf_files):
                company_name = pdf_file.name.replace('.pdf', '')
                
                temp_pdf_path = os.path.join(temp_dir, f"{company_name}.pdf")
                with open(temp_pdf_path, "wb") as f:
                    f.write(pdf_file.getvalue())
                
                status_text.text(f"Processing {company_name}... Extracting text & Querying LLM...")
                try:
                    structured_data = extract_bd_data(temp_pdf_path, api_key, company_name, model_choice)
                    
                    json_path = os.path.join(temp_dir, f"{company_name}_extracted.json")
                    with open(json_path, "w") as f:
                        json.dump(structured_data, f, indent=2)
                        
                except Exception as e:
                    st.error(f"Error during API extraction for {company_name}: {e}")
                    continue
                
                status_text.text(f"Processing {company_name}... Populating Excel...")
                output_excel_path = os.path.join(temp_dir, f"{company_name}_Filled_BD.xlsx")
                
                try:
                    template_to_use = template_path
                    wb = load_workbook(template_path)
                    sheet_name = wb.sheetnames[0]
                    
                    # Safe Expand Shareholding Table if necessary
                    required_rows = len(structured_data["fields"]["shareholding"]["rows"])
                    if required_rows > 2: 
                        status_text.text(f"Safely expanding Shareholding table...")
                        expanded_template_path = os.path.join(temp_dir, "expanded_template.xlsx")
                        
                        delta = max(0, required_rows - 2) 
                        safe_insert_rows.insert_rows_safe(
                            src_path=template_path, 
                            out_path=expanded_template_path, 
                            sheet_name=sheet_name, 
                            threshold=43, 
                            delta=delta
                        )
                        template_to_use = expanded_template_path

                    fill_factsheet.fill(
                        data_path=json_path, 
                        template_path=template_to_use, 
                        out_path=output_excel_path, 
                        sheet_name=sheet_name 
                    )
                    
                except Exception as e:
                    st.error(f"Error populating Excel for {company_name}: {e}")
                    continue
                
                with open(output_excel_path, "rb") as f:
                    final_excel_bytes = f.read()
                    
                st.download_button(
                    label=f"Download Populated BD Format ({company_name})",
                    data=final_excel_bytes,
                    file_name=f"{company_name}_Populated_BD_Format.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"dl_btn_{i}"
                )
                
                progress_bar.progress((i + 1) / len(pdf_files))
                
            status_text.text("Processing Complete!")